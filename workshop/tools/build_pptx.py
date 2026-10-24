"""Build an editable PowerPoint deck from slides.js or slides.es.js.

Usage: python build_pptx.py <slides.js> <output.pptx>
Needs python-pptx (export-slides.sh runs it with `uv run --with python-pptx`).
Image paths in the slides are relative to the workshop folder, next to slides.js.
"""
import html
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x00, 0x26, 0x78)
NAVY_DARK = RGBColor(0x00, 0x1A, 0x54)
ORANGE = RGBColor(0xFF, 0x66, 0x00)
PEACH = RGBColor(0xFF, 0xB2, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0xE6, 0xEB, 0xF3)
MUTED = RGBColor(0xBD, 0xC8, 0xD8)
GREY = RGBColor(0x94, 0xA3, 0xB8)
FONT = "Arial"

W, H = Inches(13.333), Inches(7.5)
MARGIN_X = Inches(0.9)
RIGHT_COL = Inches(3.6)


def load_slides(path):
    src = Path(path).read_text(encoding="utf-8")
    start = src.index("[", src.index("const slides"))
    end = src.index("\n];", start) + 2
    return json.loads(src[start:end])


def clean(text):
    return html.unescape(re.sub(r"<[^>]+>", "", text))


def set_bullet(paragraph, char="•"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(Emu(Inches(0.32))))
    p_pr.set("indent", str(-Emu(Inches(0.28))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for old in p_pr.findall(qn(tag)):
            p_pr.remove(old)
    bu_clr = p_pr.makeelement(qn("a:buClr"), {})
    srgb = bu_clr.makeelement(qn("a:srgbClr"), {"val": "FF6600"})
    bu_clr.append(srgb)
    p_pr.append(bu_clr)
    p_pr.append(p_pr.makeelement(qn("a:buChar"), {"char": char}))


def add_run(paragraph, text, size, color, bold=False):
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.line.fill.background()
    bg.fill.gradient()
    bg.fill.gradient_angle = 45
    stops = bg.fill.gradient_stops
    stops[0].color.rgb = NAVY
    stops[0].position = 0
    stops[1].color.rgb = NAVY_DARK
    stops[1].position = 1.0
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), H)
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE


def footer(slide, number, total, label):
    box = slide.shapes.add_textbox(MARGIN_X, H - Inches(0.6), W - 2 * MARGIN_X, Inches(0.35))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    add_run(p, f"{label}   {number} / {total}", 10, MUTED)


def body(slide, s, width):
    box = slide.shapes.add_textbox(MARGIN_X, Inches(0.7), width, H - Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0

    p = tf.paragraphs[0]
    add_run(p, clean(s["eyebrow"]).upper(), 13, PEACH, bold=True)
    p.space_after = Pt(6)

    p = tf.add_paragraph()
    add_run(p, clean(s["title"]), 44, WHITE, bold=True)
    p.space_after = Pt(18)
    p.line_spacing = 1.0

    if s.get("text"):
        p = tf.add_paragraph()
        add_run(p, clean(s["text"]), 20, TEXT)
        p.space_after = Pt(8)

    for item in s.get("bullets", []):
        p = tf.add_paragraph()
        add_run(p, clean(item), 20, TEXT)
        set_bullet(p)
        p.space_after = Pt(10)

    if s.get("presenters"):
        p = tf.add_paragraph()
        p.space_before = Pt(18)
        add_run(p, "  ·  ".join(s["presenters"]), 20, PEACH, bold=True)


def qr_slot(slide, s, base):
    size = Inches(3.0)
    left = W - MARGIN_X - size
    top = (H - size) / 2 - Inches(0.3)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, size, size)
    frame.adjustments[0] = 0.05
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    image = base / s["qr"]
    if image.exists():
        frame.line.color.rgb = WHITE
        pad = Inches(0.15)
        slide.shapes.add_picture(str(image), left + pad, top + pad, size - 2 * pad, size - 2 * pad)
    else:
        frame.line.color.rgb = PEACH
        frame.line.width = Pt(2.5)
        frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        tf = frame.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, "QR", 36, GREY, bold=True)
    cap = slide.shapes.add_textbox(left - Inches(0.3), top + size + Inches(0.12), size + Inches(0.6), Inches(0.5))
    cap.text_frame.word_wrap = True
    p = cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, clean(s.get("qrCaption", "")), 14, MUTED)
    return image.exists()


def side_image(slide, s, base):
    image = base / s["image"]
    if not image.exists():
        return
    max_side = Inches(2.4)
    pic = slide.shapes.add_picture(str(image), 0, 0)
    ratio = min(max_side / pic.width, max_side / pic.height)
    pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left = int(W - MARGIN_X - (RIGHT_COL + pic.width) / 2)
    pic.top = int((H - pic.height) / 2)


def main(src, out):
    base = Path(src).resolve().parent
    slides = load_slides(src)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    label = "KubeVirt Without Fear · Texas Linux Fest 2026"
    qr_found = None
    for i, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        background(slide)
        has_side = bool(s.get("qr") or s.get("image"))
        body(slide, s, W - 2 * MARGIN_X - (RIGHT_COL + Inches(0.3) if has_side else 0))
        if s.get("qr"):
            qr_found = qr_slot(slide, s, base)
        elif s.get("image"):
            side_image(slide, s, base)
        footer(slide, i, len(slides), label)
    prs.save(out)
    print(f"{out}: {len(slides)} slides" + ("" if qr_found is None else f", QR image {'included' if qr_found else 'missing (empty box)'}"))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
